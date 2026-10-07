"""Regenerate the bilingual downloads using the bundled Python runtime.
Dependencies: python-docx, reportlab. Source files stay editable here.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import simpleSplit

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "site" / "downloads"
OUTPUT.mkdir(parents=True, exist_ok=True)
FONT_ROOT = Path("C:/Windows/Fonts")
pdfmetrics.registerFont(TTFont("DownloadText", str(FONT_ROOT / "arial.ttf")))
pdfmetrics.registerFont(TTFont("DownloadBold", str(FONT_ROOT / "arialbd.ttf")))

def pick(language, ru, en):
    return ru if language == "ru" else en

CHECKLISTS = {
    "ads-checklist": {
        "title": ("Чек-лист запуска Google Ads", "Google Ads launch checklist"),
        "intro": ("Проверяйте пункты на реальном аккаунте и сайте. Галочка означает, что есть подтверждение проверки, а не только намерение настроить.", "Check each item against the actual account and website. A tick means verification evidence exists, not merely an intention to configure it."),
        "sections": [
            (("Цель и экономика", "Goals and economics"), [
                ("Определены заявка, качественный лид и продажа. Критерии согласованы с командой.", "Lead, qualified lead and sale are defined and agreed with the team."),
                ("Допустимый CAC и CPL рассчитаны с учётом маржи и конверсии в продажу.", "Acceptable CAC and CPL account for margin and lead-to-sale conversion."),
                ("Бюджет, валюта и условия остановки расходов подтверждены владельцем.", "Budget, currency and spending stop conditions are approved by the owner."),
            ]),
            (("Запросы и настройки", "Keywords and settings"), [
                ("Запросы сгруппированы по намерению; типы соответствия выбраны осознанно.", "Keywords are grouped by intent; match types are chosen deliberately."),
                ("Проверены минус-слова, география, языки и расписание показов.", "Negative keywords, locations, languages and schedules are checked."),
                ("Брендовый и небрендовый спрос разделены там, где это нужно для анализа.", "Brand and non-brand demand are separated where analysis requires it."),
            ]),
            (("Объявления и предложение", "Ads and the offer"), [
                ("Заголовки и описания согласованы; обещания подтверждаются страницей.", "Headlines and descriptions are approved; the page supports their claims."),
                ("Конечные URL, дополнительные ссылки и данные компании актуальны.", "Final URLs, additional links and business information are current."),
                ("Нет неподтверждённых гарантий, устаревших цен или недоступных услуг.", "No unsupported guarantees, outdated prices or unavailable services remain."),
            ]),
            (("Путь клиента", "Customer journey"), [
                ("Страница открывается на мобильном; оффер и следующий шаг понятны.", "The mobile page loads; the offer and next step are clear."),
                ("Успешная отправка формы проверена; ошибка валидации не создаёт лид.", "Successful form submission is verified; validation failure creates no lead."),
                ("Номера, мессенджеры и маршрутизация заявки проверены командой.", "Phone numbers, messengers and lead routing are verified by the team."),
            ]),
            (("Измерение", "Measurement"), [
                ("Основные конверсии соответствуют цели; микродействия учитываются отдельно.", "Primary conversions fit the goal; micro-actions are measured separately."),
                ("Одна заявка не считается несколько раз; тестовые события отделены.", "One lead is not counted multiple times; test events are separated."),
                ("В URL и параметрах аналитики нет имён, email, телефонов и сообщений.", "Analytics URLs and parameters contain no names, emails, phones or messages."),
            ]),
            (("Контроль после запуска", "Post-launch review"), [
                ("Назначены ответственный, дата проверки и окно оценки результата.", "An owner, review date and evaluation window are assigned."),
                ("Отчёт связывает расходы с качеством лидов и продажами в CRM.", "Reporting links spend to lead quality and CRM sales."),
                ("Критические ошибки устранены; план изменений и отката зафиксирован.", "Critical issues are resolved; change and rollback plans are recorded."),
            ]),
        ],
        "source": "support.google.com/google-ads/answer/6325025",
    },
    "seo-checklist": {
        "title": ("Чек-лист SEO-аудита", "SEO audit checklist"),
        "intro": ("Начните с важных страниц и типовых шаблонов. Для каждой проблемы сохраняйте пример URL, доказательство и критерий приёмки исправления.", "Start with important pages and representative templates. For each issue, record an example URL, evidence and acceptance criteria for the fix."),
        "sections": [
            (("Обход и индексация", "Crawling and indexing"), [
                ("Важные страницы доступны и возвращают ожидаемый HTTP-статус.", "Important pages are available and return the expected HTTP status."),
                ("robots.txt и noindex не закрывают нужный контент по ошибке.", "robots.txt and noindex do not accidentally block required content."),
                ("Sitemap содержит актуальные канонические URL, без ошибок и редиректов.", "The sitemap contains current canonical URLs without errors or redirects."),
            ]),
            (("Структура и намерение", "Structure and intent"), [
                ("Запросы сопоставлены со страницами и реальными задачами пользователей.", "Queries are mapped to pages and real user needs."),
                ("Дублирующие страницы выявлены; объединение или различия обоснованы.", "Duplicate pages are identified; consolidation or distinctions are justified."),
                ("Городские страницы содержат полезные различия, а не только замену названия.", "City pages offer useful differences rather than just a substituted name."),
            ]),
            (("Содержание", "Content"), [
                ("Title, основной заголовок и описание ясно объясняют содержание страницы.", "Title, main heading and description explain the page clearly."),
                ("Материал отвечает на задачу, содержит примеры и проверяемые источники.", "Content answers the task with examples and verifiable sources."),
                ("Устаревшие цены, условия, контакты и неподтверждённые обещания исправлены.", "Outdated prices, terms, contacts and unsupported claims are corrected."),
            ]),
            (("Ссылки и техническая проверка", "Links and technical review"), [
                ("Важные страницы связаны навигацией; страницы-сироты проверены.", "Important pages are linked through navigation; orphan pages are reviewed."),
                ("Нет сломанных внутренних ссылок и лишних цепочек редиректов.", "No broken internal links or unnecessary redirect chains remain."),
                ("Canonical и языковые связи соответствуют фактической структуре сайта.", "Canonicals and language relationships fit the actual site structure."),
            ]),
            (("Мобильный путь и удобство", "Mobile journey and usability"), [
                ("Нет горизонтального переполнения; текст, меню и формы доступны с телефона.", "No horizontal overflow; text, menus and forms work on a phone."),
                ("Изображения имеют разумный размер; загрузка не мешает целевому действию.", "Images are appropriately sized; loading does not obstruct the key action."),
                ("Формы, клавиатурный фокус и подписи элементов проверены.", "Forms, keyboard focus and control labels are verified."),
            ]),
            (("Измерение и план", "Measurement and planning"), [
                ("Исходные клики, показы и заявки зафиксированы с периодом и сегментами.", "Baseline clicks, impressions and leads include periods and segments."),
                ("Задачи ранжированы по влиянию и усилиям, есть ответственные и сроки.", "Tasks are ranked by impact and effort, with owners and due dates."),
                ("В отчёте учитываются сезонность, релизы и ограничения данных.", "Reporting accounts for seasonality, releases and data limitations."),
            ]),
        ],
        "source": "developers.google.com/search/docs/fundamentals/seo-starter-guide",
    },
}

def pdf_text(pdf, text, x, y, width, font="DownloadText", size=10.5, leading=15):
    pdf.setFont(font, size)
    for line in simpleSplit(text, font, size, width):
        pdf.drawString(x, y, line)
        y -= leading
    return y

def create_checklist(slug, data, language):
    filename = OUTPUT / f"{slug}-{language}.pdf"
    pdf = canvas.Canvas(str(filename), pagesize=letter)
    pdf.setTitle(pick(language, *data["title"]))
    pdf.setAuthor("Ads by Kanapiya")
    for page in range(2):
        pdf.setFillColorRGB(0, 0, 0)
        pdf_text(pdf, "ADS BY KANAPIYA", 48, 749, 500, size=8)
        y = pdf_text(pdf, pick(language, *data["title"]), 48, 712, 515, "DownloadBold", 23, 29)
        y -= 10
        y = pdf_text(pdf, pick(language, *data["intro"]) if page == 0 else pick(language, "Продолжение проверки. Критические ошибки устраняются до запуска или расширения работ.", "Continue the review. Resolve critical issues before launch or expanding the work."), 48, y, 515)
        y -= 23
        for section, items in data["sections"][page * 3:page * 3 + 3]:
            y = pdf_text(pdf, pick(language, *section), 48, y, 515, "DownloadBold", 13, 18) - 10
            for item in items:
                pdf.setStrokeColorRGB(.55, .62, .45)
                pdf.rect(49, y - 1, 10, 10, fill=0, stroke=1)
                y = pdf_text(pdf, pick(language, *item), 70, y, 490) - 13
            y -= 9
        if page == 1:
            y = pdf_text(pdf, pick(language, "Блокирующая проблема и ответственный", "Blocking issue and owner"), 48, y, 515, "DownloadBold", 11) - 8
            pdf.setStrokeColorRGB(.8, .83, .77)
            for offset in [0, 23]:
                pdf.line(48, y - offset, 564, y - offset)
        assert y > 100, f"PDF overflow in {filename.name} page {page+1}: {y}"
        pdf.setStrokeColorRGB(.82, .84, .79)
        pdf.line(48, 77, 564, 77)
        pdf_text(pdf, pick(language, "Для своих проектов. Не является гарантией результата. Версия 1.0 / 2026-10", "For your own projects. Not a performance guarantee. Version 1.0 / 2026-10"), 48, 62, 480, size=8, leading=11)
        pdf_text(pdf, "Source: " + data["source"], 48, 43, 480, size=7, leading=10)
        pdf_text(pdf, f"{page+1} / 2", 534, 62, 40, size=8)
        pdf.showPage()
    pdf.save()

BRIEFS = {
    "advertising-brief": {
        "title": ("Бриф на рекламу", "Advertising brief"),
        "intro": ("Заполняйте по фактам. Неизвестное отмечайте «нужно уточнить». Не включайте пароли, платёжные данные и личные данные клиентов.", "Use facts; mark unknowns for clarification. Do not include passwords, payment details or customers' personal information."),
        "sections": [
            (("Проект и ответственный", "Project and owner"), ("Компания, сайт, контакт для согласования и роль в проекте.", "Company, website, approval contact and project role.")),
            (("Продукт и предложение", "Product and offer"), ("Что рекламируем, цены, преимущества с подтверждением, ограничения и сезонность.", "Offer, pricing, evidence-backed advantages, limitations and seasonality.")),
            (("Аудитория и география", "Audience and locations"), ("Кто покупает, какую задачу решает, языки, города и зоны обслуживания.", "Who buys, their need, languages, cities and service areas.")),
            (("Цель и экономика", "Goal and economics"), ("Что считается лидом и продажей, средний чек, маржа, допустимые CAC и CPL.", "Lead and sale definitions, average order value, margin, acceptable CAC and CPL.")),
            (("Бюджет и сроки", "Budget and timing"), ("Валюта, рекламный бюджет отдельно от работ, период теста и ограничения расходов.", "Currency, media spend separate from fees, test period and spending limits.")),
            (("История рекламы", "Advertising history"), ("Что запускали, источники данных, сильные и слабые результаты, известные проблемы.", "Previous activity, data sources, outcomes and known problems.")),
            (("Страницы и материалы", "Pages and assets"), ("URL посадочных страниц, тексты, изображения, доказательства и ответственный за контент.", "Landing URLs, copy, images, evidence and the content owner.")),
            (("Лиды и измерение", "Leads and measurement"), ("GA4, GTM, CRM, звонки; время ответа, критерии квалификации и ответственный за продажи.", "GA4, GTM, CRM, calls; response time, qualification criteria and sales owner.")),
            (("Доступы и согласования", "Access and approvals"), ("Какие системы доступны через приглашение, кто согласует бюджет, тексты и публикацию.", "Systems accessible by invitation; who approves budget, copy and publishing.")),
            (("Приёмка и следующий шаг", "Acceptance and next step"), ("Состав работ, отчётность, условия остановки, дата проверки и открытые вопросы.", "Scope, reporting, stop conditions, review date and open questions.")),
        ],
    },
    "website-brief": {
        "title": ("Бриф на сайт и интеграции", "Website and integrations brief"),
        "intro": ("Документ помогает согласовать объём до разработки. Заполните вместе с владельцем продукта и контента. Секреты передавайте отдельно.", "Agree scope before development. Complete with the product and content owners. Share secrets separately through an agreed secure channel."),
        "sections": [
            (("Проект и бизнес задача", "Project and business goal"), ("Компания, сайт, ответственный и результат, ради которого нужен сайт.", "Company, website, owner and the business outcome the site should support.")),
            (("Аудитория и действия", "Audience and actions"), ("Кто пользуется сайтом, языки, устройства и три основных пользовательских пути.", "Audience, languages, devices and three main user journeys.")),
            (("Структура страниц", "Page structure"), ("Какие страницы, категории и типы записей нужны; что входит в первый релиз.", "Pages, categories and content types; what belongs in the first release.")),
            (("Контент и дизайн", "Content and design"), ("Кто готовит тексты, фото, логотипы; примеры по стилю и причины выбора.", "Owners of copy, photos and logos; visual references and reasons for choosing them.")),
            (("Формы и события", "Forms and events"), ("Поля, валидация, получатель заявок, согласия, успешные и ошибочные сценарии.", "Fields, validation, lead recipient, consent, success and failure scenarios.")),
            (("Платформа и редактирование", "Platform and editing"), ("CMS или стек, что редактируется без разработчика, роли и права пользователей.", "CMS or stack, content editable without development, user roles and permissions.")),
            (("Интеграции", "Integrations"), ("CRM, аналитика, оплата и другие системы; действия, источники данных и ограничения.", "CRM, analytics, payments and other systems; actions, data sources and constraints.")),
            (("Качество и доступность", "Quality and accessibility"), ("Мобильная версия, клавиатурная навигация, SEO, скорость, резервное копирование.", "Mobile layout, keyboard navigation, SEO, performance and backups.")),
            (("Миграция и запуск", "Migration and launch"), ("Домен, хостинг, старые URL, редиректы, тестовая среда и план отката.", "Domain, hosting, old URLs, redirects, staging and rollback plan.")),
            (("Приёмка и сопровождение", "Acceptance and maintenance"), ("Критерии готовности, сроки, бюджет, ответственные и что происходит после запуска.", "Acceptance criteria, timing, budget, owners and post-launch arrangements.")),
        ],
    },
}

def make_doc(title, language):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.72), Inches(.65)
    section.left_margin, section.right_margin = Inches(.75), Inches(.75)
    section.header_distance, section.footer_distance = Inches(.3), Inches(.3)
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Header", "Footer"]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing = 1.08
        style.paragraph_format.space_after = Pt(7)
        for border in list(style.element.xpath("./w:pPr/w:pBdr")):
            border.getparent().remove(border)
    doc.styles["Title"].font.size = Pt(27)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.space_after = Pt(13)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.styles["Heading 2"].font.bold = True
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    for name in ["Header", "Footer"]:
        doc.styles[name].font.size = Pt(8)
    section.header.paragraphs[0].text = "ADS BY KANAPIYA"
    footer = section.footer.paragraphs[0]
    footer.text = pick(language, "Шаблон для ваших проектов   |   Версия 1.0 / 2026-10   |   ", "Template for your projects   |   Version 1.0 / 2026-10   |   ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = title
    doc.core_properties.author = "Ads by Kanapiya"
    return doc

def answer_lines(doc, count=2):
    for _ in range(count):
        paragraph = doc.add_paragraph(" " * 3)
        paragraph.paragraph_format.space_after = Pt(9)
        paragraph.paragraph_format.line_spacing = 1
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in {"val": "single", "sz": "3", "color": "D9D9D9", "space": "3"}.items():
            bottom.set(qn("w:" + key), value)
        borders.append(bottom)
        paragraph._p.get_or_add_pPr().append(borders)

def create_brief(slug, data, language):
    title = pick(language, *data["title"])
    doc = make_doc(title, language)
    for page in range(2):
        if page:
            doc.add_page_break()
        doc.add_paragraph(title if not page else pick(language, "Детали и согласование", "Details and approval"), "Title")
        doc.add_paragraph(pick(language, *data["intro"]) if not page else pick(language, "Продолжение брифа. Фиксируйте договорённости и вопросы, которые нужно закрыть до старта.", "Continue the brief. Record agreements and questions to resolve before the start."))
        for index, (heading, prompt) in enumerate(data["sections"][page*5:page*5+5], start=page*5+1):
            doc.add_paragraph(f"{index:02d}  " + pick(language, *heading), "Heading 2")
            doc.add_paragraph(pick(language, *prompt))
            answer_lines(doc, 2)
    doc.save(OUTPUT / f"{slug}-{language}.docx")

def report_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for index, text in enumerate(headers):
        table.rows[0].cells[index].text = text
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_index])
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                margin = OxmlElement("w:"+side)
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "223548" if not row_index else ("F5F6F7" if row_index%2 else "FFFFFF"))
            props.append(shade)
            borders = OxmlElement("w:tcBorders")
            for side in ["top", "left", "bottom", "right"]:
                edge = OxmlElement("w:"+side)
                for key, value in {"val":"single", "sz":"4", "color":"D9D9D9"}.items():
                    edge.set(qn("w:"+key), value)
                borders.append(edge)
            props.append(borders)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                if col_index and len(headers)>3:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255,255,255)
    doc.add_paragraph()

def create_report(language):
    title = pick(language, "Месячный отчёт по маркетингу", "Monthly marketing report")
    doc = make_doc(title, language)
    doc.add_paragraph(title, "Title")
    doc.add_paragraph(pick(language, "Заполните реальными данными. Пустые поля не означают ноль. Сравнивайте периоды одинаковой длины с одинаковыми определениями показателей.", "Fill with actual data. Blank fields do not mean zero. Compare equal-length periods with consistent metric definitions."))
    doc.add_paragraph(pick(language, "Проект __________________  Автор __________________", "Project __________________  Author __________________"))
    doc.add_paragraph(pick(language, "Период __________  Сравнение __________  Валюта __________", "Period __________  Comparison __________  Currency __________"))
    doc.add_paragraph(pick(language, "Источники и фильтры _____________________________________", "Sources and filters _____________________________________"))
    doc.add_paragraph(pick(language, "Основные показатели", "Key metrics"), "Heading 1")
    metrics = [("Расход", "Spend"), ("Клики", "Clicks"), ("Заявки", "Leads"), ("Качественные лиды", "Qualified leads"), ("Продажи", "Sales"), ("Выручка", "Revenue"), ("CPL", "CPL"), ("Стоимость качественного лида", "Qualified CPL")]
    report_table(doc, [pick(language,"Метрика","Metric"),pick(language,"Предыдущий","Previous"),pick(language,"Текущий","Current"),pick(language,"Изменение","Change")], [[pick(language,*metric),"—","—","—"] for metric in metrics], [2.4,1.53,1.53,1.54])
    doc.add_paragraph(pick(language, "Выводы за период", "Period findings"), "Heading 1")
    doc.add_paragraph(pick(language, "Что изменилось, почему это важно и чем подтверждается вывод. Отделяйте наблюдение от гипотезы.", "What changed, why it matters and what supports the finding. Separate observations from hypotheses."))
    answer_lines(doc, 2)
    doc.add_page_break()
    doc.add_paragraph(pick(language, "Решения и следующий период", "Decisions and the next period"), "Title")
    doc.add_paragraph(pick(language, "План действий", "Action plan"), "Heading 1")
    report_table(doc, [pick(language,"Действие и гипотеза","Action and hypothesis"),pick(language,"Ответственный","Owner"),pick(language,"Срок и критерий","Due date and criterion")], [["________________","________","____________"],["________________","________","____________"],["________________","________","____________"]], [3.0,1.55,2.45])
    doc.add_paragraph(pick(language, "Определения и формулы", "Definitions and formulas"), "Heading 1")
    for ru,en in [
        ("CPL = расход / число заявок. Стоимость качественного лида = расход / число лидов, прошедших согласованную квалификацию.", "CPL = spend / leads. Qualified CPL = spend / leads meeting the agreed qualification criteria."),
        ("Изменение в процентах = (текущее - предыдущее) / предыдущее × 100. При нулевой базе показывайте абсолютную разницу, а не бесконечный рост.", "Percentage change = (current - previous) / previous × 100. With a zero baseline, report the absolute difference instead of infinite growth."),
        ("Выручка не равна прибыли. Укажите, учитываются ли налоги, возвраты, комиссии, себестоимость и расходы на работу команды.", "Revenue is not profit. State whether taxes, refunds, fees, cost of goods and team costs are included."),
    ]:
        doc.add_paragraph(pick(language,ru,en))
    doc.add_paragraph(pick(language, "Ограничения данных", "Data limitations"), "Heading 1")
    doc.add_paragraph(pick(language, "Отметьте задержку продаж, атрибуцию, сезонность, тестовый трафик, пропуски измерения и несопоставимые сегменты.", "Note sales delays, attribution, seasonality, test traffic, measurement gaps and non-comparable segments."))
    answer_lines(doc, 2)
    doc.add_paragraph(pick(language, "Проверка и согласование", "Review and approval"), "Heading 1")
    doc.add_paragraph(pick(language, "Проверил __________________  Дата __________  Следующая встреча __________", "Reviewed by __________________  Date __________  Next review __________"))
    doc.save(OUTPUT / f"monthly-report-{language}.docx")

if __name__ == "__main__":
    import sys
    mode = sys.argv[1] if len(sys.argv)>1 else "all"
    for language in ["ru", "en"]:
        if mode in ["all", "pdf"]:
            for slug, data in CHECKLISTS.items():
                create_checklist(slug, data, language)
        if mode in ["all", "docx"]:
            for slug, data in BRIEFS.items():
                create_brief(slug, data, language)
            create_report(language)
    print(f"Generated {mode} downloads in {OUTPUT}")
